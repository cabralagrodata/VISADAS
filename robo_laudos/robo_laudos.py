"""
Robô de Laudos
==============
1) APRENDER: lê o(s) vídeo(s) da pasta "Aula Laudos", transcreve o áudio,
   captura quadros da tela e usa o Claude para extrair o passo a passo e o
   modelo de laudo ensinado na aula -> salva em "instrucoes_laudo.md".
2) GERAR: para cada arquivo/pasta em "entrada", produz um laudo seguindo
   exatamente o que foi ensinado no vídeo -> salva .docx em "laudos".

Uso (dentro da pasta Aula Laudos):
    python robo_laudos.py aprender
    python robo_laudos.py gerar
"""
import base64
import subprocess
import sys
from pathlib import Path

import anthropic

PASTA = Path(__file__).resolve().parent
ENTRADA = PASTA / "entrada"
SAIDA = PASTA / "laudos"
TRABALHO = PASTA / "_robo"
INSTRUCOES = PASTA / "instrucoes_laudo.md"
MODELO = "claude-opus-5-5"
VIDEOS = {".mp4", ".mkv", ".avi", ".mov", ".wmv", ".webm", ".m4v"}
IMAGENS = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}
MAX_QUADROS = 40

cliente = anthropic.Anthropic()  # usa a variável de ambiente ANTHROPIC_API_KEY


# ---------------------------------------------------------------- utilidades
def ffmpeg():
    import imageio_ffmpeg
    return imageio_ffmpeg.get_ffmpeg_exe()


def img_bloco(caminho: Path):
    dados = base64.standard_b64encode(caminho.read_bytes()).decode()
    return {"type": "image", "source": {"type": "base64",
            "media_type": IMAGENS[caminho.suffix.lower()], "data": dados}}


def perguntar(conteudo, sistema, max_tokens=16000):
    with cliente.messages.stream(model=MODELO, max_tokens=max_tokens, system=sistema,
                                 messages=[{"role": "user", "content": conteudo}]) as s:
        return s.get_final_message().content[0].text


def ler_texto(arq: Path) -> str:
    ext = arq.suffix.lower()
    if ext in {".txt", ".md", ".csv", ".json"}:
        return arq.read_text(encoding="utf-8", errors="ignore")
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join(p.extract_text() or "" for p in PdfReader(arq).pages)
    if ext == ".docx":
        import docx
        d = docx.Document(arq)
        linhas = [p.text for p in d.paragraphs]
        for t in d.tables:
            linhas += [" | ".join(c.text for c in r.cells) for r in t.rows]
        return "\n".join(linhas)
    if ext in {".xlsx", ".xlsm"}:
        import openpyxl
        wb = openpyxl.load_workbook(arq, data_only=True)
        return "\n".join(f"[{ws.title}] " + " | ".join("" if v is None else str(v) for v in r)
                         for ws in wb for r in ws.iter_rows(values_only=True))
    return ""


# ---------------------------------------------------------------- aprender
def transcrever(video: Path) -> str:
    destino = TRABALHO / f"{video.stem}_transcricao.txt"
    if destino.exists():
        return destino.read_text(encoding="utf-8")
    print(f"  Transcrevendo {video.name} (pode demorar)...")
    from faster_whisper import WhisperModel
    modelo = WhisperModel("small", device="auto", compute_type="int8")
    segs, _ = modelo.transcribe(str(video), language="pt", vad_filter=True)
    texto = "\n".join(f"[{int(s.start)//60:02d}:{int(s.start)%60:02d}] {s.text.strip()}" for s in segs)
    destino.write_text(texto, encoding="utf-8")
    return texto


def capturar_quadros(video: Path) -> list[Path]:
    pasta = TRABALHO / f"{video.stem}_quadros"
    if not pasta.exists():
        pasta.mkdir(parents=True)
        print(f"  Capturando telas de {video.name}...")
        # só grava quadro quando a tela muda bastante (troca de slide/janela)
        subprocess.run([ffmpeg(), "-loglevel", "error", "-i", str(video),
                        "-vf", "select='gt(scene,0.08)',scale=1280:-1", "-vsync", "vfr",
                        str(pasta / "q_%04d.jpg")], check=True)
    quadros = sorted(pasta.glob("*.jpg"))
    if len(quadros) > MAX_QUADROS:
        passo = len(quadros) / MAX_QUADROS
        quadros = [quadros[int(i * passo)] for i in range(MAX_QUADROS)]
    return quadros


def aprender():
    videos = [v for v in PASTA.iterdir() if v.suffix.lower() in VIDEOS]
    if not videos:
        sys.exit(f"Nenhum vídeo encontrado em {PASTA}")
    TRABALHO.mkdir(exist_ok=True)
    conteudo = []
    for v in videos:
        print(f"Analisando {v.name}")
        conteudo.append({"type": "text", "text": f"=== TRANSCRIÇÃO DO VÍDEO {v.name} ===\n{transcrever(v)}"})
        quadros = capturar_quadros(v)
        conteudo.append({"type": "text", "text": f"=== {len(quadros)} TELAS CAPTURADAS DO VÍDEO {v.name} (em ordem) ==="})
        conteudo += [img_bloco(q) for q in quadros]
    conteudo.append({"type": "text", "text":
        "Com base na aula acima (transcrição + telas), escreva um manual completo em Markdown "
        "para que outra pessoa produza laudos EXATAMENTE como ensinado. Inclua:\n"
        "1. Tipo/finalidade do laudo e normas citadas.\n"
        "2. Dados de entrada necessários (o que precisa ser coletado/informado).\n"
        "3. Passo a passo de análise, cálculos, critérios e regras de decisão.\n"
        "4. MODELO do laudo: todas as seções na ordem certa, com textos-padrão literais mostrados "
        "na aula e marcadores {{CAMPO}} onde entram os dados.\n"
        "5. Erros comuns e cuidados mencionados pelo professor.\n"
        "Não invente o que não aparece na aula; marque como [NÃO INFORMADO NA AULA] quando faltar."})
    print("Extraindo o método de laudo com o Claude...")
    INSTRUCOES.write_text(perguntar(conteudo, "Você é um perito técnico que documenta procedimentos com precisão."),
                          encoding="utf-8")
    print(f"Pronto! Revise o manual em: {INSTRUCOES}")


# ---------------------------------------------------------------- gerar
def salvar_docx(markdown: str, destino: Path):
    import docx
    d = docx.Document()
    for linha in markdown.splitlines():
        l = linha.rstrip()
        if l.startswith("#"):
            nivel = min(len(l) - len(l.lstrip("#")), 4)
            d.add_heading(l.lstrip("# ").strip(), level=nivel)
        elif l.startswith(("- ", "* ")):
            d.add_paragraph(l[2:].replace("**", ""), style="List Bullet")
        elif l.startswith("|") and set(l.replace("|", "").strip()) <= set("-: "):
            continue
        elif l:
            d.add_paragraph(l.replace("**", ""))
    d.save(destino)


def gerar():
    if not INSTRUCOES.exists():
        sys.exit("Rode primeiro: python robo_laudos.py aprender")
    ENTRADA.mkdir(exist_ok=True)
    SAIDA.mkdir(exist_ok=True)
    casos = sorted(ENTRADA.iterdir())
    if not casos:
        sys.exit(f"Coloque os dados de cada laudo em {ENTRADA} (um arquivo ou uma subpasta por laudo).")
    manual = INSTRUCOES.read_text(encoding="utf-8")
    for caso in casos:
        arquivos = sorted(caso.rglob("*")) if caso.is_dir() else [caso]
        conteudo = [{"type": "text", "text": f"=== MANUAL DO LAUDO (aprendido na aula) ===\n{manual}\n\n=== DADOS DO CASO: {caso.name} ==="}]
        for a in arquivos:
            if a.suffix.lower() in IMAGENS:
                conteudo += [{"type": "text", "text": f"[imagem {a.name}]"}, img_bloco(a)]
            elif txt := ler_texto(a):
                conteudo.append({"type": "text", "text": f"[arquivo {a.name}]\n{txt}"})
        conteudo.append({"type": "text", "text":
            "Elabore o laudo final deste caso seguindo rigorosamente o manual (mesmas seções, "
            "textos-padrão, cálculos e critérios). Use somente os dados fornecidos; onde faltar "
            "informação escreva [PENDENTE: o que falta]. Responda apenas com o laudo em Markdown."})
        print(f"Gerando laudo: {caso.name}")
        laudo = perguntar(conteudo, "Você é um perito que redige laudos técnicos precisos e formais.")
        (SAIDA / f"Laudo_{caso.stem}.md").write_text(laudo, encoding="utf-8")
        salvar_docx(laudo, SAIDA / f"Laudo_{caso.stem}.docx")
    print(f"Laudos salvos em: {SAIDA}")


if __name__ == "__main__":
    {"aprender": aprender, "gerar": gerar}.get(sys.argv[1] if len(sys.argv) > 1 else "", lambda: print(__doc__))()
